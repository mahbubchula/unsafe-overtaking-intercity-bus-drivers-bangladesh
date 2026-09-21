#!/usr/bin/env python3
"""Validate questionnaire code inventories without changing source documents."""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree


ROOT = Path(__file__).resolve().parents[1]
QUESTIONNAIRES = {
    "English DOCX": ROOT / "questionnaire/english/questionnaire-en.docx",
    "English PDF": ROOT / "questionnaire/english/questionnaire-en.pdf",
    "Bangla DOCX": ROOT / "questionnaire/bangla/questionnaire-bn.docx",
    "Bangla PDF": ROOT / "questionnaire/bangla/questionnaire-bn.pdf",
}
CODEBOOK = ROOT / "documentation/codebook/master-codebook.xlsx"
REQUIRED_CODES = {
    "A1", "A2", "A3", "A4", "A5", "A6", "D1",
    "C2", "D2", "D3", "E4", "F1",
    "UO1", "UO2", "UO3", "UO4", "UO5", "S1", "S2",
}
CODE_RE = re.compile(r"(?<![A-Z0-9])(?:UO|[A-Z]{1,3})\d+(?![A-Z0-9])")


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        names = [
            name for name in archive.namelist()
            if name == "word/document.xml" or name.startswith("word/header")
            or name.startswith("word/footer")
        ]
        chunks: list[str] = []
        for name in names:
            root = ElementTree.fromstring(archive.read(name))
            chunks.extend(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
        return "\n".join(chunks)


def docx_paragraph_texts(path: Path) -> list[str]:
    """Return document-body paragraph text in XML order, including table cells."""
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    paragraphs: list[str] = []
    for paragraph in (node for node in root.iter() if node.tag.endswith("}p")):
        text = "".join(
            node.text or "" for node in paragraph.iter() if node.tag.endswith("}t")
        ).strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required to inspect PDF files") from exc
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def extract_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        return docx_text(path)
    if path.suffix.lower() == ".pdf":
        return pdf_text(path)
    raise ValueError(f"Unsupported format: {path.suffix}")


def ordered_unique_codes(text: str) -> list[str]:
    return list(dict.fromkeys(match.group(0).upper() for match in CODE_RE.finditer(text.upper())))


def codebook_inventory(path: Path) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to inspect the master codebook") from exc
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook["Variable register"]
        inventory = [
            str(row[0]).strip().upper()
            for row in worksheet.iter_rows(min_row=2, values_only=True)
            if row[0] is not None and CODE_RE.fullmatch(str(row[0]).strip().upper())
        ]
    finally:
        workbook.close()
    return list(dict.fromkeys(inventory))


def administered_order(path: Path, expected: set[str]) -> list[str]:
    """Find first-use item order, starting at the first A1 question block."""
    paragraphs = docx_paragraph_texts(path)
    start = next(
        (index for index, text in enumerate(paragraphs) if re.match(r"^A1(?:\s|$)", text)),
        None,
    )
    if start is None:
        return []
    found: list[str] = []
    for text in paragraphs[start:]:
        match = re.match(r"^((?:UO|[A-Z]{1,3})\d+)(?:\s|$)", text.upper())
        if match and match.group(1) in expected and match.group(1) not in found:
            found.append(match.group(1))
            if len(found) == len(expected):
                break
    return found


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--allow-missing-files",
        action="store_true",
        help="Report absent source documents as warnings during repository preparation.",
    )
    args = parser.parse_args()
    errors: list[str] = []
    warnings: list[str] = []
    inventories: dict[str, list[str]] = {}

    if not CODEBOOK.exists():
        errors.append(f"Master codebook is missing: {CODEBOOK.relative_to(ROOT)}")
        expected_codes = set(REQUIRED_CODES)
    else:
        try:
            expected_inventory = codebook_inventory(CODEBOOK)
        except Exception as exc:
            errors.append(f"Could not read master codebook inventory: {exc}")
            expected_inventory = []
        if len(expected_inventory) != 45:
            errors.append(
                f"Master codebook contains {len(expected_inventory)} unique item codes; expected 45"
            )
        expected_codes = set(expected_inventory) or set(REQUIRED_CODES)
        missing_required = sorted(REQUIRED_CODES - expected_codes)
        if missing_required:
            errors.append(
                "Master codebook is missing required codes: " + ", ".join(missing_required)
            )

    for label, path in QUESTIONNAIRES.items():
        if not path.exists():
            message = f"{label} is missing: {path.relative_to(ROOT)}"
            (warnings if args.allow_missing_files else errors).append(message)
            continue
        if path.stat().st_size == 0:
            errors.append(f"{label} is empty: {path.relative_to(ROOT)}")
            continue
        try:
            text = extract_text(path)
        except Exception as exc:  # diagnostic tool: preserve the original file
            errors.append(f"Could not read {label}: {exc}")
            continue
        all_codes = set(ordered_unique_codes(text))
        codes = sorted(all_codes & expected_codes)
        inventories[label] = codes
        if len(codes) != len(expected_codes):
            errors.append(
                f"{label} contains {len(codes)} of {len(expected_codes)} codebook item codes"
            )
        missing = sorted(expected_codes - set(codes))
        if missing:
            errors.append(f"{label} is missing codebook item codes: {', '.join(missing)}")
        if path.suffix.lower() == ".docx":
            order = administered_order(path, expected_codes)
            if "A6" not in order or "D1" not in order:
                errors.append(f"{label} administered-question order could not locate A6 and D1")
            elif order.index("D1") != order.index("A6") + 1:
                errors.append(f"{label} does not place D1 immediately after A6")
        if label == "English DOCX":
            normalized = re.sub(r"\s+", " ", text).lower()
            if "overtaking-related near-crash" not in normalized:
                errors.append("English DOCX does not contain the required term 'overtaking-related near-crash'")

    for left, right in (("English DOCX", "English PDF"), ("Bangla DOCX", "Bangla PDF")):
        if left in inventories and right in inventories and set(inventories[left]) != set(inventories[right]):
            errors.append(f"Code inventory differs between {left} and {right}")
    if "English DOCX" in inventories and "Bangla DOCX" in inventories:
        en = set(inventories["English DOCX"])
        bn = set(inventories["Bangla DOCX"])
        if en != bn:
            errors.append(
                "English and Bangla DOCX code inventories differ: "
                f"English-only={sorted(en - bn)}, Bangla-only={sorted(bn - en)}"
            )

    for message in warnings:
        print(f"WARNING: {message}")
    for message in errors:
        print(f"ERROR: {message}", file=sys.stderr)
    print(f"Questionnaire verification: {len(errors)} error(s), {len(warnings)} warning(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
